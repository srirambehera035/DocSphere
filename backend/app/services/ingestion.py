import io
import base64
from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from PIL import Image
import httpx
from app.core.config import settings

class IngestionError(Exception):
    pass

class DocumentIngestionService:
    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".png", ".jpg", ".jpeg", ".webp"}
    MAGIC_SIGNATURES = {
        ".pdf": [b"%PDF-"],
        ".docx": [b"PK\x03\x04"],
        ".png": [b"\x89PNG\r\n\x1a\n"],
        ".jpg": [b"\xff\xd8\xff"],
        ".jpeg": [b"\xff\xd8\xff"],
        ".webp": [b"RIFF"]
    }

    def validate_file(self, filename: str, content_type: str, file_size: int, file_bytes: bytes, max_size: int) -> None:
        if file_size > max_size:
            raise IngestionError(f"File size ({file_size} bytes) exceeds maximum limit ({max_size} bytes)")
        if file_size < 4:
            raise IngestionError("Uploaded file is empty or corrupted")

        ext = Path(filename).suffix.lower()
        if ext not in self.ALLOWED_EXTENSIONS:
            raise IngestionError(f"Unsupported file extension: {ext}. Allowed formats: {', '.join(sorted(self.ALLOWED_EXTENSIONS))}")

        if ext in self.MAGIC_SIGNATURES:
            valid_sig = any(file_bytes.startswith(sig) for sig in self.MAGIC_SIGNATURES[ext])
            if not valid_sig and ext != ".webp":
                raise IngestionError(f"File signature mismatch for {ext}. File content does not match extension header.")

    def extract_text_from_pdf(self, file_bytes: bytes) -> tuple[str, list[tuple[int, str]], int]:
        stream = io.BytesIO(file_bytes)
        reader = PdfReader(stream)
        pages_text: list[tuple[int, str]] = []
        total_pages = len(reader.pages)

        for idx, page in enumerate(reader.pages):
            page_no = idx + 1
            text = page.extract_text() or ""
            cleaned = text.strip()
            pages_text.append((page_no, cleaned))

        full_text = "\n\n".join(
            f"[PAGE {p_no}]\n{p_text}" for p_no, p_text in pages_text if p_text
        )

        has_scanned_pages = any(len(t.strip()) < 40 for _, t in pages_text)
        if (not full_text.strip() or has_scanned_pages) and settings.GEMINI_API_KEY:
            try:
                ocr_text = self.gemini_vision_ocr(file_bytes, "application/pdf")
                if ocr_text.strip():
                    full_text = ocr_text
                    pages_text = [(1, ocr_text)]
            except Exception:
                pass

        return full_text, pages_text, max(1, total_pages)

    def extract_text_from_docx(self, file_bytes: bytes) -> tuple[str, list[tuple[int, str]], int]:
        stream = io.BytesIO(file_bytes)
        doc = DocxDocument(stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        full_text = "[PAGE 1]\n" + "\n\n".join(paragraphs)
        return full_text, [(1, "\n\n".join(paragraphs))], 1

    def gemini_vision_ocr(self, file_bytes: bytes, mime_type: str) -> str:
        if not settings.GEMINI_API_KEY:
            return ""
        b64_data = base64.b64encode(file_bytes).decode("utf-8")
        url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key={settings.GEMINI_API_KEY}"
        payload = {
            "contents": [{
                "parts": [
                    {"text": "Perform high precision OCR on this document. Extract all textual content verbatim preserving layout and page sections. Prefix sections with [PAGE 1], [PAGE 2] etc."},
                    {
                        "inlineData": {
                            "mimeType": mime_type,
                            "data": b64_data
                        }
                    }
                ]
            }],
            "generationConfig": {
                "temperature": 0.0
            }
        }
        with httpx.Client(timeout=45.0) as client:
            res = client.post(url, json=payload)
            if res.status_code == 200:
                data = res.json()
                candidates = data.get("candidates", [])
                if candidates:
                    parts = candidates[0].get("content", {}).get("parts", [])
                    if parts:
                        return parts[0].get("text", "")
        return ""

    def extract_text_from_image(self, file_bytes: bytes, ext: str) -> tuple[str, list[tuple[int, str]], int]:
        mime = "image/png" if ext == ".png" else "image/jpeg"
        if settings.GEMINI_API_KEY:
            try:
                ocr_text = self.gemini_vision_ocr(file_bytes, mime)
                if ocr_text.strip():
                    return f"[PAGE 1]\n{ocr_text}", [(1, ocr_text)], 1
            except Exception:
                pass

        try:
            import pytesseract
            image = Image.open(io.BytesIO(file_bytes))
            text = pytesseract.image_to_string(image)
            return f"[PAGE 1]\n{text.strip()}", [(1, text.strip())], 1
        except Exception:
            return "[PAGE 1]\n[Image document ingested. OCR text extraction requires Gemini API key or Tesseract binary.]", [(1, "")], 1

    def extract(self, filename: str, file_bytes: bytes) -> tuple[str, list[tuple[int, str]], int]:
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return self.extract_text_from_pdf(file_bytes)
        elif ext in {".docx", ".doc"}:
            return self.extract_text_from_docx(file_bytes)
        elif ext in {".png", ".jpg", ".jpeg", ".webp"}:
            return self.extract_text_from_image(file_bytes, ext)
        else:
            raise IngestionError(f"Unsupported file format: {ext}")

ingestion_service = DocumentIngestionService()
